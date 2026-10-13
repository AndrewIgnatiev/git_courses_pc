from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.responses import FileResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from docx import Document
from docx.shared import Pt
from datetime import datetime
import json
import os
import tempfile
import shutil

app = FastAPI(title="ИПР Generator")

# CORS — чтобы браузер мог обращаться
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # для локалки можно *
    allow_methods=["*"],
    allow_headers=["*"],
)

TEMPLATE_PATH = "test_case/shablon_ipr.docx"


class IprRequest(BaseModel):
    employee: dict
    section_1_results: list
    section_2_goals: list
    section_3_actions: list


# ============================================================
# Утилиты (те же, что в fill_ipr.py)
# ============================================================
def fill_cell_multiline(cell, text, size=9):
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    for r in p.runs:
        r.text = ""
    while len(p.runs) > 1:
        p.runs[-1]._element.getparent().remove(p.runs[-1]._element)
    lines = str(text).split("\n") if text else [""]
    if not p.runs:
        run = p.add_run(lines[0])
    else:
        p.runs[0].text = lines[0]
        run = p.runs[0]
    run.font.size = Pt(size)
    for line in lines[1:]:
        new_p = cell.add_paragraph()
        new_run = new_p.add_run(line)
        new_run.font.size = Pt(size)


def split_methods(methods_str):
    return [m.strip() for m in str(methods_str).split(",") if m.strip()]


def get_num_rows(a):
    n_actions = len(a.get("actions", []))
    n_methods = len(split_methods(a.get("methods", "")))
    return max(n_actions, n_methods, 1)


def format_goal_with_smart(goal):
    text = goal.get("goal", "")
    smart = goal.get("smart", {})
    if smart:
        parts = [f"{L}: {smart[L]}" for L in ['S', 'M', 'A', 'R', 'T'] if smart.get(L)]
        if parts:
            text += "\n" + "\n".join(parts)
    return text


# ============================================================
# Генерация
# ============================================================
@app.post("/generate")
async def generate_ipr(req: IprRequest):
    if not os.path.exists(TEMPLATE_PATH):
        raise HTTPException(500, f"Шаблон не найден: {TEMPLATE_PATH}")

    doc = Document(TEMPLATE_PATH)
    data = req.dict()

    emp = data.get("employee", {})
    name = emp.get("name", "Сотрудник")
    today = datetime.now().strftime("%d.%m.%Y")

    # Заголовок
    p1 = doc.paragraphs[1]
    for run in p1.runs:
        if "<ФИО>" in run.text:
            run.text = run.text.replace("_____<ФИО>______.", f"{name}.")
        elif "_____<ФИО>______" in run.text:
            run.text = run.text.replace("_____<ФИО>______", name)

    p2 = doc.paragraphs[2]
    for run in p2.runs:
        if "______________" in run.text:
            run.text = run.text.replace("______________", f": {today}")

    p3 = doc.paragraphs[3]
    for run in p3.runs:
        if "______________" in run.text:
            run.text = run.text.replace(" ______________", f": {today}")

    # Раздел 1
    results = data.get("section_1_results", [])
    p6 = doc.paragraphs[6]
    for run in p6.runs:
        run.text = ""

    result_paragraphs = [doc.paragraphs[8], doc.paragraphs[9], doc.paragraphs[10]]
    for i, p in enumerate(result_paragraphs):
        if i < len(results):
            for run in p.runs:
                run.text = ""
            p.add_run(f"{i+1}. {results[i]}")

    # Раздел 2 — Цели
    goals = data.get("section_2_goals", [])
    goals_table = doc.tables[0]
    while len(goals_table.rows) > 1:
        goals_table._element.remove(goals_table.rows[-1]._element)
    for _ in range(len(goals)):
        goals_table.add_row()

    for i, g in enumerate(goals):
        row = goals_table.rows[i + 1]
        deadline = g.get("deadline", "") or g.get("smart", {}).get("T", "")
        fill_cell_multiline(row.cells[0], str(g.get("id", i + 1)), size=9)
        fill_cell_multiline(row.cells[1], format_goal_with_smart(g), size=9)
        fill_cell_multiline(row.cells[2], g.get("comment", ""), size=9)
        fill_cell_multiline(row.cells[3], deadline, size=9)

    # Раздел 3 — Действия
    actions = data.get("section_3_actions", [])
    actions_table = doc.tables[1]
    while len(actions_table.rows) > 1:
        actions_table._element.remove(actions_table.rows[-1]._element)

    total_rows = sum(get_num_rows(a) for a in actions)
    for _ in range(total_rows):
        actions_table.add_row()

    row_cursor = 1
    for a in actions:
        n = get_num_rows(a)
        methods = split_methods(a.get("methods", ""))
        acts = a.get("actions", [])

        for col in [0, 3, 4, 5]:
            if n > 1:
                try:
                    top = actions_table.rows[row_cursor].cells[col]
                    bottom = actions_table.rows[row_cursor + n - 1].cells[col]
                    top.merge(bottom)
                except Exception:
                    pass

        comp_text = a.get("competency", "")
        if a.get("indicator"):
            comp_text += f"\n({a['indicator']})"

        fill_cell_multiline(actions_table.rows[row_cursor].cells[0], comp_text, size=9)
        fill_cell_multiline(actions_table.rows[row_cursor].cells[3], a.get("who_helps", ""), size=9)
        fill_cell_multiline(actions_table.rows[row_cursor].cells[4], a.get("deadline", ""), size=9)
        fill_cell_multiline(actions_table.rows[row_cursor].cells[5], str(a.get("goal_id", "")), size=9)

        for j in range(n):
            m = methods[j] if j < len(methods) else ""
            act = acts[j] if j < len(acts) else ""
            fill_cell_multiline(actions_table.rows[row_cursor + j].cells[1], m, size=9)
            fill_cell_multiline(actions_table.rows[row_cursor + j].cells[2], f"• {act}" if act else "", size=9)

        row_cursor += n

    # Сохраняем во временный файл
    safe_name = name.replace(" ", "_").replace("/", "-")
    filename = f"ИПР_{safe_name}_{datetime.now().strftime('%Y-%m-%d')}.docx"
    tmp_dir = tempfile.mkdtemp()
    output_path = os.path.join(tmp_dir, filename)
    doc.save(output_path)

    return FileResponse(
        output_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=filename
    )


@app.get("/health")
async def health():
    return {"status": "ok", "template_exists": os.path.exists(TEMPLATE_PATH)}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8000)