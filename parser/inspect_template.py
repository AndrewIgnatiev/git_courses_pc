from docx import Document

doc = Document("test_case/shablon_ipr.docx")

print("=== ПАРАГРАФЫ ===")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"[{i}] style={p.style.name}")
        print(f"     text: {p.text[:150]}")
        for j, run in enumerate(p.runs):
            if run.text.strip():
                print(f"       run[{j}]: '{run.text[:80]}'")

print("\n=== ТАБЛИЦЫ ===")
for i, table in enumerate(doc.tables):
    print(f"\nТаблица {i}: {len(table.rows)} строк × {len(table.columns)} колонок")
    for j, row in enumerate(table.rows):
        cells = [c.text[:40].replace('\n', '⏎') for c in row.cells]
        print(f"  Row {j}: {cells}")

print("\n=== КАРТИНКИ ===")
print(f"Inline shapes: {len(doc.inline_shapes)}")
for i, shape in enumerate(doc.inline_shapes):
    print(f"  Shape {i}: {shape.type}, {shape.width}×{shape.height}")