from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.responses import FileResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from docx import Document
from docx.shared import Pt
from datetime import datetime
from pathlib import Path
import json
import os
import sys
import subprocess
import tempfile
import shutil

app = FastAPI(title="ИПР Generator")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# === Пути ===
PARSER_DIR = Path(__file__).parent                # parser/
ROOT_DIR = PARSER_DIR.parent                       # корень проекта
TEST_DIR = PARSER_DIR / "test_case"
TEMPLATE_PATH = TEST_DIR / "shablon_ipr.docx"
TEST_DIR.mkdir(exist_ok=True)

# Имена файлов, которые ждут парсеры
FILE_MAP = {
    "eos": "soft_matrica.xlsx",
    "hard": "hard_skills_matrix.xlsx",
    "transcript": "transcript_BOR.docx",
}

PIPELINE_STEPS = [
    ("parse_matrix.py", "Парсинг матрицы софтов"),
    ("parse_hard.py", "Парсинг чек-листа хардов"),
    ("build_prompt.py", "Формирование промпта"),
    ("call_llm.py", "Генерация через LLM"),
]


class IprRequest(BaseModel):
    employee: dict
    section_1_results: list
    section_2_goals: list
    section_3_actions: list


# ============================================================
# Утилиты заполнения шаблона
# ============================================================
def fill_cell_multiline(cell, text, size=9):
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    for r in p.runs:
        r.text = ""
    while len(p.runs) > 1:
        p.runs[-1]._element.getparent().remove(p.runs[-1]._element)
    lines = str(text).split("\n") if text else [""]
    if not p.runs:
        run = p.add_run(lines[0])
    else:
        p.runs[0].text = lines[0]
        run = p.runs[0]
    run.font.size = Pt(size)
    for line in lines[1:]:
        new_p = cell.add_paragraph()
        new_run = new_p.add_run(line)
        new_run.font.size = Pt(size)


def split_methods(methods_str):
    return [m.strip() for m in str(methods_str).split(",") if m.strip()]


def format_goal_with_smart(goal):
    text = goal.get("goal", "")
    smart = goal.get("smart", {})
    if smart:
        parts = [f"{L}: {smart[L]}" for L in ['S', 'M', 'A', 'R', 'T'] if smart.get(L)]
        if parts:
            text += "\n" + "\n".join(parts)
    return text


def generate_word_from_json(data: dict) -> str:
    """Заполняет shablon_ipr.docx данными из JSON, возвращает путь к готовому файлу."""
    if not os.path.exists(TEMPLATE_PATH):
        raise HTTPException(500, f"Шаблон не найден: {TEMPLATE_PATH}")

    doc = Document(TEMPLATE_PATH)

    emp = data.get("employee", {})
    name = emp.get("name", "Сотрудник")
    today = datetime.now().strftime("%d.%m.%Y")

    # Заголовок
    p1 = doc.paragraphs[1]
    for run in p1.runs:
        if "<ФИО>" in run.text:
            run.text = run.text.replace("_____<ФИО>______.", f"{name}.")
        elif "_____<ФИО>______" in run.text:
            run.text = run.text.replace("_____<ФИО>______", name)

    p2 = doc.paragraphs[2]
    for run in p2.runs:
        if "______________" in run.text:
            run.text = run.text.replace("______________", f": {today}")

    p3 = doc.paragraphs[3]
    for run in p3.runs:
        if "______________" in run.text:
            run.text = run.text.replace(" ______________", f": {today}")

    # Раздел 1
    results = data.get("section_1_results", [])
    p6 = doc.paragraphs[6]
    for run in p6.runs:
        run.text = ""

    result_paragraphs = [doc.paragraphs[8], doc.paragraphs[9], doc.paragraphs[10]]
    for i, p in enumerate(result_paragraphs):
        if i < len(results):
            for run in p.runs:
                run.text = ""
            p.add_run(f"{i+1}. {results[i]}")

    # Раздел 2 — Цели
    goals = data.get("section_2_goals", [])
    goals_table = doc.tables[0]
    while len(goals_table.rows) > 1:
        goals_table._element.remove(goals_table.rows[-1]._element)
    for _ in range(len(goals)):
        goals_table.add_row()

    for i, g in enumerate(goals):
        row = goals_table.rows[i + 1]
        deadline = g.get("deadline", "") or g.get("smart", {}).get("T", "")
        fill_cell_multiline(row.cells[0], str(g.get("id", i + 1)), size=9)
        fill_cell_multiline(row.cells[1], format_goal_with_smart(g), size=9)
        fill_cell_multiline(row.cells[2], g.get("comment", ""), size=9)
        fill_cell_multiline(row.cells[3], deadline, size=9)

    # Раздел 3 — Действия (методы и actions — независимыми списками)
    actions = data.get("section_3_actions", [])
    actions_table = doc.tables[1]
    while len(actions_table.rows) > 1:
        actions_table._element.remove(actions_table.rows[-1]._element)

    for _ in range(len(actions)):
        actions_table.add_row()

    for idx, a in enumerate(actions):
        row = actions_table.rows[idx + 1]

        comp_text = a.get("competency", "")
        if a.get("indicator"):
            comp_text += f"\n({a['indicator']})"
        fill_cell_multiline(row.cells[0], comp_text, size=9)

        methods = split_methods(a.get("methods", ""))
        methods_text = "\n".join([f"• {m}" for m in methods]) if methods else ""
        fill_cell_multiline(row.cells[1], methods_text, size=9)

        acts = a.get("actions", [])
        actions_text = "\n".join([f"• {x}" for x in acts]) if acts else ""
        fill_cell_multiline(row.cells[2], actions_text, size=9)

        fill_cell_multiline(row.cells[3], a.get("who_helps", ""), size=9)
        fill_cell_multiline(row.cells[4], a.get("deadline", ""), size=9)
        fill_cell_multiline(row.cells[5], str(a.get("goal_id", "")), size=9)

    safe_name = name.replace(" ", "_").replace("/", "-")
    filename = f"ИПР_{safe_name}_{datetime.now().strftime('%Y-%m-%d')}.docx"
    tmp_dir = tempfile.mkdtemp()
    output_path = os.path.join(tmp_dir, filename)
    doc.save(output_path)
    return output_path


# ============================================================
# Endpoint 1: JSON → Word
# ============================================================
@app.post("/generate")
async def generate_ipr(req: IprRequest):
    path = generate_word_from_json(req.dict())
    return FileResponse(
        path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=os.path.basename(path),
    )


# ============================================================
# Endpoint 2: 3 файла → JSON (для заполнения интерфейса)
# ============================================================
@app.post("/generate-from-files")
async def generate_from_files(
    eos: UploadFile = File(...),
    hard: UploadFile = File(...),
    transcript: UploadFile = File(...),
):
    # 1. Сохраняем файлы под именами, которые ждут парсеры
    try:
        with open(TEST_DIR / FILE_MAP["eos"], "wb") as f:
            f.write(await eos.read())
        with open(TEST_DIR / FILE_MAP["hard"], "wb") as f:
            f.write(await hard.read())
        with open(TEST_DIR / FILE_MAP["transcript"], "wb") as f:
            f.write(await transcript.read())
    except Exception as e:
        raise HTTPException(500, f"Ошибка сохранения файлов: {e}")

    # 2. Запускаем пайплайн
    logs = []
    for script, label in PIPELINE_STEPS:
        logs.append(f"▶ {label} ({script})")
        try:
            result = subprocess.run(
                [sys.executable, script],
                cwd=PARSER_DIR,
                capture_output=True,
                text=True,
                encoding="utf-8",
                timeout=900,
            )
        except subprocess.TimeoutExpired:
            raise HTTPException(500, f"Шаг '{script}' завис (>15 мин). Логи:\n" + "\n".join(logs))
        if result.returncode != 0:
            raise HTTPException(500, f"Ошибка в {script}.\n{result.stderr[-800:]}\n\nЛоги:\n" + "\n".join(logs))
        logs.append(result.stdout[-1000:] if result.stdout else "")

    # 3. Читаем ipr_draft.json
    draft_path = PARSER_DIR / "ipr_draft.json"
    if not draft_path.exists():
        raise HTTPException(500, "Не создан ipr_draft.json. Логи:\n" + "\n".join(logs))

    with open(draft_path, encoding="utf-8") as f:
        data = json.load(f)

    # 4. Возвращаем JSON — фронт заполнит интерфейс
    return data


# ============================================================
# Раздача HTML и статики из корня
# ============================================================
@app.get("/")
async def index():
    return FileResponse(ROOT_DIR / "ipr-builder.html")


@app.get("/catalog.js")
async def catalog_js():
    return FileResponse(ROOT_DIR / "catalog.js")


@app.get("/catalog-courses.js")
async def catalog_courses_js():
    return FileResponse(ROOT_DIR / "catalog-courses.js")


@app.get("/health")
async def health():
    return {"status": "ok", "template_exists": os.path.exists(TEMPLATE_PATH)}


if __name__ == "__main__":
    import uvicorn
    print("=" * 60)
    print("  Конструктор ИПР запущен")
    print("  Открой в браузере: http://localhost:8000")
    print("=" * 60)
    uvicorn.run(app, host="127.0.0.1", port=8000)