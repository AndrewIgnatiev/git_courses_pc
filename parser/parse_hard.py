import openpyxl
import json
import sys

FILE = "test_case/hard_skills_matrix.xlsx"

# === Открываем книгу и берём первый лист ===
try:
    wb = openpyxl.load_workbook(FILE, data_only=True)
except Exception as e:
    print(f"❌ Не удалось открыть файл «{FILE}»: {e}")
    sys.exit(1)

if not wb.sheetnames:
    print(f"❌ В файле «{FILE}» нет ни одного листа.")
    sys.exit(1)

first_sheet_name = wb.sheetnames[0]
ws = wb[first_sheet_name]
print(f"📄 Читаем первый лист: «{first_sheet_name}» ({ws.max_row} строк × {ws.max_column} колонок)")

# === Собираем все строки ===
rows = []
for r in range(1, ws.max_row + 1):
    rows.append({
        "r": r,
        "a": ws.cell(row=r, column=1).value,
        "b": ws.cell(row=r, column=2).value,
        "c": ws.cell(row=r, column=3).value,
        "d": ws.cell(row=r, column=4).value,
    })


def next_significant(idx):
    """Следующая строка, у которой есть текст в C или число в B."""
    for j in range(idx + 1, len(rows)):
        row = rows[j]
        if row["c"] is not None or (isinstance(row["b"], int) and row["b"] in (2, 3, 4, 5)):
            return row
    return None


SKIP_PREFIXES = ("Другое:", "Зоны роста:", "!После", "Количество респондентов", "% - от общего")
SKIP_EXACT = ("Сотрудник", "Я")

competencies = []
current_category = None
current_competency = None
current_items = []

for i, row in enumerate(rows):
    b, c, d = row["b"], row["c"], row["d"]

    # Подпункт: B = число 2-5 и C = описание
    if isinstance(b, int) and b in (2, 3, 4, 5) and c:
        current_items.append({
            "level": b,
            "description": str(c).strip()[:200],
            "checked": bool(d) if d is not None else False,
        })
        continue

    # Название: C заполнено, B пусто
    if c and b is None:
        c_str = str(c).strip()

        if any(c_str.startswith(p) for p in SKIP_PREFIXES):
            continue
        if c_str in SKIP_EXACT:
            continue
        if len(c_str) < 2:
            continue

        nxt = next_significant(i)
        if nxt is None:
            continue

        next_has_level = isinstance(nxt["b"], int) and nxt["b"] in (2, 3, 4, 5)

        if next_has_level:
            # Это КОМПЕТЕНЦИЯ
            if current_competency and current_items:
                competencies.append({
                    "category": current_category,
                    "competency": current_competency,
                    "items": current_items,
                })
            current_competency = c_str
            current_items = []
        else:
            # Это КАТЕГОРИЯ
            if current_competency and current_items:
                competencies.append({
                    "category": current_category,
                    "competency": current_competency,
                    "items": current_items,
                })
                current_competency = None
                current_items = []
            current_category = c_str

# Последняя компетенция
if current_competency and current_items:
    competencies.append({
        "category": current_category,
        "competency": current_competency,
        "items": current_items,
    })

# === Считаем уровни ===
result = []
warnings = []
for comp in competencies:
    achieved = [item["level"] for item in comp["items"] if item["checked"]]
    max_level = max(achieved) if achieved else 0

    if len(achieved) > 1:
        warnings.append(
            f"«{comp['competency']}»: отмечено {len(achieved)} уровней {sorted(achieved)}. "
            f"Берём максимальный: {max_level}."
        )

    result.append({
        "category": comp["category"],
        "competency": comp["competency"],
        "max_level": max_level,
        "achieved_levels": sorted(achieved),
    })

# === Проверка результата ===
if not result:
    print("❌ Не найдено ни одной компетенции в файле.")
    print("   Проверь структуру: столбец B — уровни 2-5, столбец C — названия.")
    sys.exit(1)

no_level = [c for c in result if c["max_level"] == 0]
if no_level:
    print(f"\n⚠️  {len(no_level)} компетенций без отмеченных уровней (max_level = 0):")
    for c in no_level[:5]:
        print(f"   • {c['competency']}")
    if len(no_level) > 5:
        print(f"   ... и ещё {len(no_level) - 5}")

if warnings:
    print(f"\n⚠️  Обнаружено {len(warnings)} компетенций с несколькими галочками:")
    for w in warnings[:10]:
        print(f"   • {w}")
    if len(warnings) > 10:
        print(f"   ... и ещё {len(warnings) - 10}")

# === Сохраняем ===
with open("hard_skills.json", "w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False, indent=2)

print(f"\n✅ Всего компетенций: {len(result)}")
print(f"   С отмеченными уровнями: {len(result) - len(no_level)}")
print(f"   Сохранено в hard_skills.json")

# Показываем превью
print("\n=== Превью ===")
for r in result[:10]:
    print(f"  [{r['max_level']}] {r['category']} → {r['competency']}")
if len(result) > 10:
    print(f"  ... и ещё {len(result) - 10}")