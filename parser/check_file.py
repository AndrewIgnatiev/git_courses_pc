import openpyxl

FILE = "test_case/soft_matrica.xlsx"
wb = openpyxl.load_workbook(FILE, data_only=True)

print("=== Коллеги_% — полная структура ===")
ws = wb["Коллеги_%"]
for r in range(1, 155):
    a = ws.cell(row=r, column=1).value
    b = ws.cell(row=r, column=2).value
    c = ws.cell(row=r, column=3).value
    d = ws.cell(row=r, column=4).value
    e = ws.cell(row=r, column=5).value
    # Печатаем только непустые строки
    if any([a, b, c, d, e]):
        # Форматируем числа
        def fmt(v):
            if v is None: return "—"
            if isinstance(v, float): return f"{v:.3f}"
            return str(v)[:60]
        print(f"  R{r}: A={fmt(a)} | B={fmt(b)} | C={fmt(c)} | D={fmt(d)} | E={fmt(e)}")

print("\n=== Оценка коллег — заголовки (строки 1-3) ===")
ws = wb["Оценка коллег"]
# Заголовки в первой строке
for r in range(1, 4):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 120)]
    non_empty = [(i+1, v) for i, v in enumerate(vals) if v is not None]
    print(f"\nСтрока {r} (непустых {len(non_empty)}):")
    for col_idx, v in non_empty[:20]:  # первые 20 столбцов
        print(f"  Col {col_idx}: {str(v)[:70]}")