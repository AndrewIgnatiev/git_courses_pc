import json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# === Читаем данные ===
with open("ipr_draft.json", encoding="utf-8") as f:
    ipr = json.load(f)

emp = ipr["employee"]

# === Создаём документ ===
doc = Document()

# Поля страницы
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

# Шрифт по умолчанию
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)

# === ЗАГОЛОВОК ===
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run(f"Личный план развития {emp.get('name', '—')}")
run.bold = True
run.font.size = Pt(15)

# Подзаголовок
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run(
    f"Дата проведения Беседы о развитии: {emp.get('bor_date', '25.09.2026')}\n"
    f"Дата утверждения плана о развитии: {emp.get('plan_date', '20.10.2026')}"
).italic = True

doc.add_paragraph()

# === РАЗДЕЛ 1: РЕЗУЛЬТАТЫ БОР ===
h1 = doc.add_heading("1. Результаты Беседы о развитии", level=2)
for r in h1.runs:
    r.font.color.rgb = RGBColor(0x21, 0x2F, 0x45)
    r.font.name = "Times New Roman"

for i, item in enumerate(ipr.get("section_1_results", []), 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.add_run(f"{i}. ").bold = True
    p.add_run(item)

doc.add_paragraph()

# === РАЗДЕЛ 2: ЦЕЛИ ===
h2 = doc.add_heading("2. Цели на следующий период", level=2)
for r in h2.runs:
    r.font.color.rgb = RGBColor(0x21, 0x2F, 0x45)
    r.font.name = "Times New Roman"

# Таблица целей: 4 колонки
goals_table = doc.add_table(rows=1, cols=4)
goals_table.style = "Table Grid"
goals_table.alignment = WD_TABLE_ALIGNMENT.CENTER

# Заголовки
hdr = goals_table.rows[0].cells
headers = ["№", "Цель", "Комментарий", "Срок достижения"]
for i, text in enumerate(headers):
    hdr[i].text = ""
    p = hdr[i].paragraphs[0]
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(10)

# Дефолтные даты для целей (если модель не заполнила)
DEFAULT_DEADLINES = ["2027-10-12", "2027-04-12", "2027-07-12"]

# Данные
for i, goal in enumerate(ipr.get("section_2_goals", [])):
    row = goals_table.add_row().cells
    row[0].text = str(goal.get("id", i + 1))
    row[1].text = goal.get("goal", "—")
    row[2].text = goal.get("comment", "")

    # Срок: если есть — берём, иначе подставляем по индексу
    deadline = goal.get("deadline")
    if not deadline or deadline == "—":
        deadline = DEFAULT_DEADLINES[i] if i < len(DEFAULT_DEADLINES) else "2027-10-12"
    row[3].text = deadline

    for cell in row:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(10)

# Ширина колонок
widths = [Cm(1), Cm(6), Cm(7), Cm(3)]
for row in goals_table.rows:
    for i, w in enumerate(widths):
        row.cells[i].width = w

doc.add_paragraph()

# === РАЗДЕЛ 3: РАЗВИВАЮЩИЕ ДЕЙСТВИЯ ===
h3 = doc.add_heading("3. Развивающие действия", level=2)
for r in h3.runs:
    r.font.color.rgb = RGBColor(0x21, 0x2F, 0x45)
    r.font.name = "Times New Roman"

# Таблица действий: 6 колонок
actions_table = doc.add_table(rows=1, cols=6)
actions_table.style = "Table Grid"
actions_table.alignment = WD_TABLE_ALIGNMENT.CENTER

# Заголовки
hdr = actions_table.rows[0].cells
headers = [
    "Компетенция / Индикатор",
    "Методы развития",
    "Конкретные действия",
    "Кто помогает",
    "Срок",
    "№ цели",
]
for i, text in enumerate(headers):
    hdr[i].text = ""
    p = hdr[i].paragraphs[0]
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(9)

# Данные
for action in ipr.get("section_3_actions", []):
    row = actions_table.add_row().cells

    # Компетенция + индикатор
    row[0].text = ""
    p0 = row[0].paragraphs[0]
    p0.add_run(action.get("competency", "—")).bold = True
    p0.add_run(f"\n({action.get('indicator', '—')})").italic = True

    # Методы
    row[1].text = action.get("methods", "—")

    # Действия — списком
    row[2].text = ""
    p2 = row[2].paragraphs[0]
    for j, a in enumerate(action.get("actions", [])):
        if j > 0:
            p2 = row[2].add_paragraph()
        p2.add_run(f"• {a}")
        p2.paragraph_format.space_after = Pt(2)

    # Кто помогает
    row[3].text = action.get("who_helps", "—")

    # Срок
    row[4].text = action.get("deadline", "—")

    # Номер цели
    row[5].text = str(action.get("goal_id", "—"))

    # Мелкий шрифт
    for cell in row:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)

# Ширина колонок
widths_actions = [Cm(3.5), Cm(3), Cm(5), Cm(2.5), Cm(2), Cm(1.5)]
for row in actions_table.rows:
    for i, w in enumerate(widths_actions):
        row.cells[i].width = w

# === Подвал ===
doc.add_paragraph()
footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Согласовано с руководителем: _________________________").italic = True

# === Сохраняем ===
name = emp.get('name', 'Сотрудник').replace(' ', '_')
output_file = f"ИПР_{name}.docx"
doc.save(output_file)

print(f"✅ ИПР сохранён: {output_file}")
print(f"   Сотрудник: {emp.get('name', '—')}")
print(f"   Должность: {emp.get('position', '—')}")
print(f"   Грейд: {emp.get('current_grade', '?')} → {emp.get('target_grade', '?')}")
print(f"   Целей: {len(ipr.get('section_2_goals', []))}")
print(f"   Компетенций в развитии: {len(ipr.get('section_3_actions', []))}")